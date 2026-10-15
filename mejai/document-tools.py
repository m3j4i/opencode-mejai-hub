import sys, json, io, os
from pathlib import Path
from xml.sax.saxutils import escape

def run(name, a):
    p = Path(a['path']).resolve()
    if name == 'audit_pdf':
        import fitz, uuid
        if p.stat().st_size > 150_000_000:
            raise ValueError('PDF exceeds 150 MB; use a dedicated incremental audit.')
        previews = Path(__file__).resolve().parent / 'dados' / 'previews' / uuid.uuid4().hex
        previews.mkdir(parents=True, exist_ok=True)
        with fitz.open(p) as doc:
            if doc.needs_pass:
                raise ValueError('An authorized unlocked PDF is required.')
            if not len(doc):
                raise ValueError('The PDF has no pages.')
            samples = sorted(set([0, min(1, len(doc)-1), len(doc)//2, max(0, len(doc)-2), len(doc)-1]))
            sections = a.get('expected_sections', [])
            found = set()
            counts = {'characters': 0, 'hebrew_characters': 0, 'replacement_characters': 0, 'empty_text_pages': 0}
            images = []
            for index, page in enumerate(doc):
                text = page.get_text()
                counts['characters'] += len(text)
                counts['hebrew_characters'] += sum('\u0590' <= c <= '\u05ff' for c in text)
                counts['replacement_characters'] += text.count('\ufffd')
                counts['empty_text_pages'] += int(not text.strip())
                found.update(section for section in sections if section.casefold() in text.casefold())
                if index in samples:
                    target = previews / f'page-{index+1}.png'
                    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(target)
                    images.append({'page': index+1, 'path': str(target), 'width_pt': page.rect.width, 'height_pt': page.rect.height})
            return {'path': str(p), 'pages': len(doc), **counts, 'missing_section_labels': [s for s in sections if s not in found], 'minimum_page_check': None if 'minimum_pages' not in a else len(doc) >= a['minimum_pages'], 'preview_images': images, 'visual_review': 'required: open the preview images and inspect legibility, fonts, direction and clipping', 'completeness': 'not proven by page count or section labels; compare full coverage against the requested source'}
    if name == 'read_document':
        ext = p.suffix.lower()
        start, limit = a.get('start', 0), a.get('limit', 20)
        if p.stat().st_size > 100_000_000:
            raise ValueError('File exceeds 100 MB; use a dedicated incremental processing script.')
        if ext == '.pdf':
            from pypdf import PdfReader
            r = PdfReader(p)
            if r.is_encrypted: raise ValueError('An authorized unlocked PDF is required.')
            total = len(r.pages)
            items = [{'page': i+1, 'text': r.pages[i].extract_text() or ''} for i in range(start,min(total,start+limit))]
        elif ext == '.docx':
            from docx import Document
            r = Document(p)
            values = [x.text for x in r.paragraphs] + [' | '.join(c.text for c in row.cells) for t in r.tables for row in t.rows]
            total, items = len(values), values[start:start+limit]
        elif ext == '.pptx':
            from pptx import Presentation
            slides = Presentation(p).slides
            total = len(slides)
            items = [{'slide':i+1,'text':'\n'.join(s.text for s in slides[i].shapes if s.has_text_frame)} for i in range(start,min(total,start+limit))]
        elif ext == '.xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(p,read_only=True,data_only=False)
            try:
                ws = wb[a['sheet']] if a.get('sheet') else wb.active
                total = ws.max_row or 0
                items = list(ws.iter_rows(min_row=start+1,max_row=min(total,start+limit),max_col=min(ws.max_column or 1,100),values_only=True)) if start<total else []
                return {'path':str(p),'sheet':ws.title,'sheets':wb.sheetnames,'rows':items,'total':total,'next':start+limit if start+limit<total else None,'note':'Formulas are not recalculated. Reading is limited to 100 columns.'}
            finally: wb.close()
        else: raise ValueError('Supported reading formats: PDF DOCX XLSX PPTX.')
        return {'path':str(p),'items':items,'total':total,'next':start+limit if start+limit<total else None,'note':'Digital PDF: extracted text without OCR. DOCX: paragraphs followed by tables.'}
    if name == 'image_info':
        from PIL import Image
        with Image.open(p) as im:
            result={'format':im.format,'width':im.width,'height':im.height,'mode':im.mode,'frames':getattr(im,'n_frames',1)}
            im.verify()
            return result
    if p.exists(): raise FileExistsError('Destination already exists. Choose another name; existing files are preserved.')
    data=io.BytesIO()
    if name == 'create_document':
        ext=p.suffix.lower()
        if ext=='.docx':
            from docx import Document
            doc=Document();doc.add_heading(a['title'],0)
            for paragraph in a['paragraphs']: doc.add_paragraph(paragraph)
            doc.save(data)
        elif ext=='.pdf':
            from reportlab.platypus import SimpleDocTemplate, Paragraph
            from reportlab.lib.styles import getSampleStyleSheet
            styles=getSampleStyleSheet()
            SimpleDocTemplate(data).build([Paragraph(escape(a['title']),styles['Title'])]+[Paragraph(escape(t).replace('\n','<br/>'),styles['BodyText']) for t in a['paragraphs']])
        else: raise ValueError('Destination must end in .docx ou .pdf.')
    elif name == 'create_spreadsheet':
        if p.suffix.lower()!='.xlsx': raise ValueError('Destination must end in .xlsx.')
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb=Workbook();ws=wb.active;ws.title=a.get('sheet','Dados')
        for row in a['rows']:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value,str): cell.data_type='s'
        for cell in ws[1]: cell.font=Font(bold=True)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(data)
    elif name == 'create_slides':
        if p.suffix.lower()!='.pptx': raise ValueError('Destination must end in .pptx.')
        from pptx import Presentation
        pr=Presentation()
        for item in a['slides']:
            slide=pr.slides.add_slide(pr.slide_layouts[1]);slide.shapes.title.text=item['title'];slide.placeholders[1].text=item['text']
        pr.save(data)
    elif name == 'merge_pdf':
        if p.suffix.lower()!='.pdf': raise ValueError('Destination must end in .pdf.')
        from pypdf import PdfWriter
        writer=PdfWriter()
        for src in a['sources']: writer.append(src)
        writer.write(data);writer.close()
    elif name == 'convert_image':
        from PIL import Image, ImageOps
        formats={'.png':'PNG','.jpg':'JPEG','.jpeg':'JPEG','.webp':'WEBP'}
        fmt=formats.get(p.suffix.lower())
        if not fmt: raise ValueError('Output must be PNG, JPEG or WEBP.')
        with Image.open(a['source']) as src:
            if getattr(src,'n_frames',1)>1: raise ValueError('Animated or multipage image: use a dedicated script to preserve frames.')
            im=ImageOps.exif_transpose(src)
            if a.get('max_width') or a.get('max_height'): im.thumbnail((a.get('max_width',im.width),a.get('max_height',im.height)))
            if fmt=='JPEG':
                rgba=im.convert('RGBA');im=Image.new('RGB',rgba.size,'white');im.paste(rgba,mask=rgba.getchannel('A'))
            im.save(data,format=fmt)
    else: raise ValueError('Unknown tool.')
    p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('xb') as out: out.write(data.getvalue())
    return {'path':str(p),'bytes':p.stat().st_size,'created':True}

if __name__=='__main__':
    try:
        request=json.load(sys.stdin)
        result=run(request['name'],request['arguments'])
        print(json.dumps(result,ensure_ascii=False,default=str))
    except Exception as e:
        print(json.dumps({'error':str(e)},ensure_ascii=False));sys.exit(1)
